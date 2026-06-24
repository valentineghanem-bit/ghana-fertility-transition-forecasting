"""md_to_docx.py — minimal Markdown -> .docx converter for the Project 15 manuscript.
Handles: #/##/###/#### headings, paragraphs, **bold**/*italic* runs, markdown tables,
> blockquotes, and - bullet lists. No external deps beyond python-docx."""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

NUMERIC = re.compile(r'^[\s\d.,×%/\-–−\[\]()<>±]+$')


def _border(cell, edge, sz):
    tcPr = cell._tc.get_or_add_tcPr()
    b = tcPr.find(qn('w:tcBorders'))
    if b is None:
        b = OxmlElement('w:tcBorders'); tcPr.append(b)
    e = b.find(qn(f'w:{edge}'))
    if e is None:
        e = OxmlElement(f'w:{edge}'); b.append(e)
    e.set(qn('w:val'), 'single'); e.set(qn('w:sz'), str(sz))
    e.set(qn('w:space'), '0'); e.set(qn('w:color'), '000000')


def add_runs(par, text):
    # split on **bold** and *italic*
    for tok in re.split(r'(\*\*[^*]+\*\*|\*[^*]+\*)', text):
        if not tok:
            continue
        if tok.startswith('**') and tok.endswith('**'):
            par.add_run(tok[2:-2]).bold = True
        elif tok.startswith('*') and tok.endswith('*'):
            par.add_run(tok[1:-1]).italic = True
        else:
            par.add_run(tok)


def convert(md_path, docx_path):
    lines = open(md_path, encoding='utf-8').read().split('\n')
    doc = Document()
    doc.styles['Normal'].font.name = 'Calibri'
    doc.styles['Normal'].font.size = Pt(11)
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if not ln.strip():
            i += 1
            continue
        if ln.startswith('|') and i + 1 < len(lines) and re.match(r'\|[-:\s|]+\|', lines[i + 1]):
            # table block
            block = []
            while i < len(lines) and lines[i].lstrip().startswith('|'):
                block.append(lines[i]); i += 1
            rows = [[c.strip() for c in r.strip().strip('|').split('|')] for r in block if not re.match(r'\|[-:\s|]+\|', r)]
            if rows:
                ncol = len(rows[0])
                tbl = doc.add_table(rows=len(rows), cols=ncol)  # no style -> borderless base
                tbl.autofit = True
                last = len(rows) - 1
                # column is numeric if all data cells are numeric -> right-align that column
                numeric_col = [all(NUMERIC.match((rows[ri][ci] or 'x').replace('**', '')) for ri in range(1, len(rows)))
                               for ci in range(ncol)]
                for ri, row in enumerate(rows):
                    for ci, cell in enumerate(row):
                        if ci >= ncol:
                            continue
                        tc = tbl.rows[ri].cells[ci]
                        p = tc.paragraphs[0]; p.text = ''
                        # strip markdown bold in data cells (Q1 tables: emphasis via ordering, not bold)
                        txt = re.sub(r'<br/?>', ' ', cell)
                        if ri > 0:
                            txt = txt.replace('**', '')
                        add_runs(p, txt)
                        if ri == 0:
                            for run in p.runs:
                                run.bold = True
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if numeric_col[ci] else WD_ALIGN_PARAGRAPH.LEFT
                        elif numeric_col[ci]:
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
                        # three-line (booktabs) borders
                        if ri == 0:
                            _border(tc, 'top', 12); _border(tc, 'bottom', 6)
                        if ri == last:
                            _border(tc, 'bottom', 12)
                doc.add_paragraph()
            continue
        mimg = re.match(r'^!\[[^\]]*\]\(([^)]+)\)\s*$', ln)
        if mimg:
            rel = mimg.group(1)
            path = rel if os.path.isabs(rel) else os.path.normpath(os.path.join(os.path.dirname(md_path), rel))
            if os.path.exists(path):
                doc.add_picture(path, width=Inches(6.0))
                doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                doc.add_paragraph(f"[missing image: {rel}]")
            i += 1
            continue
        if ln.startswith('#### '):
            doc.add_heading(ln[5:], level=4)
        elif ln.startswith('### '):
            doc.add_heading(ln[4:], level=3)
        elif ln.startswith('## '):
            doc.add_heading(ln[3:], level=2)
        elif ln.startswith('# '):
            h = doc.add_heading('', level=0); add_runs(h, ln[2:])
        elif ln.startswith('> '):
            p = doc.add_paragraph(); p.style = 'Intense Quote'; add_runs(p, ln[2:])
        elif re.match(r'^[-*] ', ln):
            p = doc.add_paragraph(style='List Bullet'); add_runs(p, ln[2:])
        elif re.match(r'^---+$', ln):
            doc.add_paragraph()
        else:
            p = doc.add_paragraph(); add_runs(p, ln)
        i += 1
    doc.save(docx_path)
    return len(doc.paragraphs), len(doc.tables)


if __name__ == '__main__':
    src = sys.argv[1] if len(sys.argv) > 1 else 'manuscript/Manuscript_Fertility_Transition_Ghana.md'
    out = sys.argv[2] if len(sys.argv) > 2 else 'manuscript/Manuscript_Fertility_Transition_Ghana.docx'
    p, t = convert(src, out)
    print(f"wrote {out}: {p} paragraphs, {t} tables")
